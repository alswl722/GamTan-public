"""회계 Excel(data/*.xlsx) 읽기 — 배출계수 · 월별단가 · 비정형 표현/라벨.

회계 담당이 Excel만 관리하면 개발자가 init_db·생성기를 실행할 때 반영된다.
Excel/시트 없거나 파싱 실패 시 예외를 던져 호출부가 하드코딩으로 폴백한다.
"""
import os

import openpyxl

DEFAULT_XLSX = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "data",
    "iM-Bridge_데이터준비_샘플.xlsx",
)

# 시트마다 연료 표기가 흔들려 정규화 (계수·단가·분류가 같은 이름을 쓰도록)
FUEL_ALIAS = {
    "도시가스(LNG등)": "도시가스",
    "도시가스(LNG)": "도시가스",
    "도시가스(lng)": "도시가스",
}


def _norm_fuel(name) -> str:
    n = (str(name).strip() if name is not None else "")
    return FUEL_ALIAS.get(n, n)


def _scope_int(s):
    """'Scope 1' -> 1, '제외' -> None."""
    if s is None:
        return None
    t = str(s)
    for d in ("1", "2", "3"):
        if d in t:
            return int(d)
    return None


def _truthy(v) -> bool:
    return str(v).strip().upper() in ("TRUE", "Y", "1", "O", "예")


def _load(path: str):
    if not os.path.exists(path):
        raise FileNotFoundError(f"회계 Excel 없음: {path}")
    return openpyxl.load_workbook(path, data_only=True)


def _sheet(wb, name: str):
    if name not in wb.sheetnames:
        raise LookupError(f"시트 '{name}' 없음 ({wb.sheetnames})")
    return wb[name]


def _rows_as_dicts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        out.append({hdr[i]: (r[i] if i < len(r) else None) for i in range(len(hdr))})
    return out


def load_emission_factors(path: str = DEFAULT_XLSX) -> list[dict]:
    """`배출계수` 시트 → EmissionFactor 입력. 계수 없는(불명/없음) 행은 스킵."""
    rows = _rows_as_dicts(_sheet(_load(path), "배출계수"))
    out = []
    for r in rows:
        gwp = r.get("배출계수_kgCO2e_per_unit")
        scope = _scope_int(r.get("Scope"))
        if gwp is None or scope is None:
            continue
        out.append(
            dict(
                fuel_type=_norm_fuel(r.get("연료/에너지")),
                scope=scope,
                category=None,
                factor_co2=float(gwp),
                factor_ch4=0.0,
                factor_n2o=0.0,
                gwp_co2e=float(gwp),
                unit=r.get("단위"),
                source=r.get("출처URL"),
            )
        )
    return out


def _price_unit(price_unit_label) -> str:
    """'원/L' -> 'L' (UnitPrice.unit은 물량 단위만 저장)."""
    s = str(price_unit_label) if price_unit_label is not None else ""
    return s.split("/", 1)[-1].strip() if "/" in s else s.strip()


def load_unit_prices(path: str = DEFAULT_XLSX, year: int = 2024) -> list[dict]:
    """환산단가 → UnitPrice 입력 (연료×12개월).

    `월별단가` 시트가 있으면 그걸 우선 사용 (회계가 월별 실제값을 넣으면 자동 반영).
    없으면 `배출계수` 시트의 단일 단가(월별_환산단가_공급가액)를 12개월에 동일
    적용한다 — 유가 월별 변동을 반영하지 못하는 데모 임시값이며, 회계 담당이
    `월별단가` 시트를 추가하는 즉시 대체된다.
    """
    wb = _load(path)
    if "월별단가" in wb.sheetnames:
        rows = _rows_as_dicts(_sheet(wb, "월별단가"))
        out = []
        for r in rows:
            price = r.get("단가")
            if price is None:
                continue
            out.append(
                dict(
                    fuel_type=_norm_fuel(r.get("연료")),
                    year=int(r.get("연")),
                    month=int(r.get("월")),
                    unit_price_krw=float(price),
                    unit=r.get("단위"),
                    source=r.get("출처"),
                )
            )
        return out

    rows = _rows_as_dicts(_sheet(wb, "배출계수"))
    out = []
    for r in rows:
        price = r.get("월별_환산단가_공급가액")
        if price is None:
            continue
        fuel = _norm_fuel(r.get("연료/에너지"))
        unit = _price_unit(r.get("단가단위"))
        source = r.get("출처URL")
        note = f"{source} (배출계수 시트 단일값, 12개월 동일 적용 — 데모 임시값)" if source else "데모 임시값 (월별 미분리)"
        for month in range(1, 13):
            out.append(
                dict(
                    fuel_type=fuel,
                    year=year,
                    month=month,
                    unit_price_krw=float(price),
                    unit=unit,
                    source=note,
                )
            )
    return out


def load_expressions(path: str = DEFAULT_XLSX) -> dict:
    """`전표_샘플` 시트 → 생성기 표현 사전.

    반환: {연료: {"exprs":[품목명…], "label":{scope,category,fuel}, "ambiguous":{HITL 품목명}}}
    """
    rows = _rows_as_dicts(_sheet(_load(path), "전표_샘플"))
    groups: dict = {}
    for r in rows:
        fuel = _norm_fuel(r.get("정답연료"))
        expr = (str(r.get("품목명")).strip() if r.get("품목명") else "")
        if not fuel or not expr:
            continue
        g = groups.setdefault(
            fuel,
            {
                "exprs": [],
                "label": {
                    "scope": _scope_int(r.get("정답Scope")),
                    "category": r.get("정답분류"),
                    "fuel": fuel,
                },
                "ambiguous": set(),
            },
        )
        if expr not in g["exprs"]:
            g["exprs"].append(expr)
        if _truthy(r.get("사람검토필요")):
            g["ambiguous"].add(expr)
    return groups
